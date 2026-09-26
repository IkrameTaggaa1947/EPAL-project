"""Build criteria_table_LURH.xlsx in the canonical criteria model.

Same 7 sheets and same columns as criteria_table_TUA.xlsx — both builders share
extraction/criteria_model.py — so the two files append cleanly in Power BI.

Where the LURH data comes from
------------------------------
Unlike TUA, LURH had no licence-level criteria table at all: the workbook
CriteriosConformidade_geral_*.xlsx is a criteria DICTIONARY (which rules exist),
with no link to individual licences. So the fact table is derived here from
`master_lurh.xlsx` -> sheet 'Conditions', which already carries, per licence and
parameter: Regime, VLE, VLE (% mín. remoção), Legislação aplicável and
Avaliação da Conformidade Legal.

Criteria flags are assigned in two passes, and the provenance of every rule is
recorded in 'Origem dos critérios':
  1. inherit the validated flags of a matching CriteriosConformidade_GLOBAL row
     (same catalogue parameter + same/equivalent conformity text);
  2. otherwise classify the Avaliação text with the SAME keyword rules already
     used by sync_criterios.py, and mark the rule 'AUTO — rever'.

Careful — CriteriosConformidade_GLOBAL's 'VLE' column is a 0/1 FLAG, not the
numeric limit. In the canonical model the numeric limit is 'VLE' and the flag is
'VLE (flag)'. Confusing the two is the single easiest way to corrupt this model.

Run:  python build_criteria_table.py [--dry-run] [--out path.xlsx]
"""
import argparse
import os
import re
import sys

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'Common_Code_PowerBI'))
import criteria_model as cm      # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
MASTER = os.getenv('EPAL_LURH_MASTER', os.path.join(HERE, 'master_lurh.xlsx'))
CRITERIA = os.getenv('EPAL_CRITERIA', os.path.join(
    HERE, 'CriteriosConformidade_geral_20250926_funcionalidades_critérios.xlsx'))
DEFAULT_OUT = os.path.abspath(os.path.join(HERE, '..', '..', 'data', 'powerbi', 'criteria_table_LURH.xlsx'))

GLOBAL_SHEET = 'CriteriosConformidade_GLOBAL'
NOTE_RE = re.compile(r'^\s*(?:\((?P<p>[a-z0-9])\)\s*)+', re.I)   # '(b)(*) Óleos e Gorduras'


# ------------------------------------------------- keyword classification
def classify(txt):
    """The keyword rules of sync_criterios.py, mapped onto the canonical flags.

    Kept deliberately identical to sync_criterios.classify() so the automatic
    part of this file and the automatic part of the criteria workbook never
    disagree. If you change one, change the other.
    """
    t = cm.norm(txt)
    c1_mensal = 'media mensal' in t
    c5_anual = 'media anual' in t
    c2_dobro = 'dobro do vle' in t
    c3_borlas = '152/97' in t and ('amostras nao conformes' in t
                                   or 'minimo anual de amostras' in t
                                   or 'quadro n' in t
                                   or 'alinea d' in t)
    c4_100 = '100%' in t or 'mais de 100' in t
    c4_150 = '150%' in t
    c6_ordem = 'ordem de grandeza' in t
    c7_gama = bool(re.search(r'5[.,]0\s*-?\s*10[.,]0', t)) or 'nenhuma amostra' in t
    art69 = ('artigo 69' in t or 'art. 69' in t
             or ('artigo n' in t and '69' in t)) and '152/97' not in t
    return (
        1,                                   # VLE (flag) — every row carries a VLE
        0,                                   # % mín remoção — set by the caller
        1 if c7_gama else 0,                 # Gama de valores (intervalo)
        1 if (c2_dobro or c4_100 or art69) else 0,   # ≤ 100% VLE (dobro)
        1 if c4_150 else 0,                  # ≤ 150% VLE
        1 if c6_ordem else 0,                # ≤ uma ordem de grandeza do VLE
        1 if c1_mensal else 0,               # média mensal ≤ VLE
        1 if c5_anual else 0,                # média anual ≤ VLE
        1 if c3_borlas else 0,               # Quadro III DL 152/97 (borlas)
    )


# ------------------------------------------------------- GLOBAL catalogue
def read_global(path):
    """CriteriosConformidade_GLOBAL -> [(catalogue param, folded note, flags, empresa)].

    Column order in that sheet is already the canonical flag order:
    VLE, % mín Remoção, Gama, ≤100%, ≤150%, ≤ordem, média mensal, média anual,
    Quadro III — indices 3..11.
    """
    if not os.path.exists(path):
        print(f'! {os.path.basename(path)} not found — every rule will be classified '
              f'automatically', file=sys.stderr)
        return []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if GLOBAL_SHEET not in wb.sheetnames:
        wb.close()
        return []
    rows = list(wb[GLOBAL_SHEET].iter_rows(values_only=True))
    wb.close()
    out = []
    for r in rows[1:]:
        if not r or not r[2]:
            continue
        flags = tuple(1 if str(v).strip() in ('1', '1.0') else 0 for v in r[3:12])
        note = cm.norm(r[12] if len(r) > 12 else '')
        out.append((cm.param_catalogo(r[2]), note, flags, r[1]))
    return out


def match_global(param_cat, aval_text, catalogue):
    """Inherit validated flags when the same rule is already described in GLOBAL.

    Matching is on the folded conformity text: exact first, then containment
    (GLOBAL stores an ABBREVIATED note, so it is usually a substring of the full
    licence text). A 40-character floor keeps short generic notes from matching
    everything.
    """
    a = cm.norm(aval_text)
    if not a:
        return None
    same_param = [g for g in catalogue if g[0] == param_cat and g[1]]
    for _, note, flags, emp in same_param:
        if note == a:
            return flags, emp
    for _, note, flags, emp in sorted(same_param, key=lambda g: -len(g[1])):
        if len(note) >= 40 and (note in a or a in note):
            return flags, emp
    return None


# ------------------------------------------------------------ master read
def read_master(path):
    """master_lurh.xlsx -> canonical records (one per Conditions row)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for s in ('Conditions', 'Licenses'):
        if s not in wb.sheetnames:
            raise SystemExit(f"'{s}' not found in {os.path.basename(path)}")

    lic_rows = list(wb['Licenses'].iter_rows(values_only=True))
    li = {h: i for i, h in enumerate(lic_rows[0])}
    licences = {}
    for r in lic_rows[1:]:
        key = r[li['Nº Licença']]
        if key:
            licences[key] = {
                'Estabelecimento': r[li['Estabelecimento']],
                'Empresa': r[li['Requerente']] if 'Requerente' in li else None,
                'Ficheiro fonte': r[li['file_name']],
            }

    cond_rows = list(wb['Conditions'].iter_rows(values_only=True))
    wb.close()
    ci = {h: i for i, h in enumerate(cond_rows[0])}
    need = ['Nº Licença', 'Regime', 'Parâmetro', 'VLE', 'VLE (% mín. remoção)',
            'Legislação aplicável', 'Avaliação da Conformidade Legal']
    missing = [c for c in need if c not in ci]
    if missing:
        raise SystemExit(f"master 'Conditions' is missing columns: {missing}")

    recs = []
    for r in cond_rows[1:]:
        key = r[ci['Nº Licença']]
        if not key or not r[ci['Parâmetro']]:
            continue
        raw_param = str(r[ci['Parâmetro']])
        note = NOTE_RE.match(raw_param)
        recs.append({
            'IdLicenca': key,
            'Estabelecimento': licences.get(key, {}).get('Estabelecimento'),
            'Empresa': licences.get(key, {}).get('Empresa'),
            'Ficheiro fonte': licences.get(key, {}).get('Ficheiro fonte'),
            'Regime': r[ci['Regime']] or 'Normal',
            'Parâmetro (documento)': raw_param,
            'Parâmetro (catálogo)': cm.param_catalogo(raw_param),
            'VLE': r[ci['VLE']],
            'VLE (% mín. remoção)': r[ci['VLE (% mín. remoção)']],
            # LURH keeps footnote letters in the Legislação cell / on the parameter
            'Nota/Código': (note.group(0).strip() if note else None),
            'Interpretação → critérios': None,      # filled below
            'Texto da avaliação de conformidade (lei)': r[ci['Avaliação da Conformidade Legal']],
            'Legislação aplicável (texto)': r[ci['Legislação aplicável']],
            'Empresa (catálogo)': None,
            'Origem dos critérios': None,
            'flags': None,
        })
    return recs


def assign_flags(recs, catalogue):
    """Fill flags / Empresa (catálogo) / Origem for every record."""
    stats = {cm.ORIG_GLOBAL: 0, cm.ORIG_AUTO: 0}
    for r in recs:
        hit = match_global(r['Parâmetro (catálogo)'],
                           r['Texto da avaliação de conformidade (lei)'], catalogue)
        if hit:
            flags, empresa = hit
            r['Origem dos critérios'] = cm.ORIG_GLOBAL
            r['Empresa (catálogo)'] = empresa
            r['Interpretação → critérios'] = (
                'Critérios herdados de CriteriosConformidade_GLOBAL (regra já validada)')
        else:
            flags = classify(r['Texto da avaliação de conformidade (lei)'])
            r['Origem dos critérios'] = cm.ORIG_AUTO
            r['Interpretação → critérios'] = (
                'Critérios classificados automaticamente a partir do texto da avaliação '
                '(mesmas regras de sync_criterios.py) — POR VALIDAR')
        # a minimum-removal percentage in the licence sets its own flag,
        # whatever the text says
        flags = list(flags)
        if str(r['VLE (% mín. remoção)'] or '').strip():
            flags[1] = 1
        if not str(r['VLE'] or '').strip():
            flags[0] = 0
        r['flags'] = tuple(flags)
        stats[r['Origem dos critérios']] += 1
    return stats


def main(argv):
    ap = argparse.ArgumentParser()
    ap.add_argument('--master', default=MASTER)
    ap.add_argument('--criteria', default=CRITERIA)
    ap.add_argument('--out', dest='dst', default=DEFAULT_OUT)
    ap.add_argument('--dry-run', action='store_true')
    a = ap.parse_args(argv)

    recs = read_master(a.master)
    catalogue = read_global(a.criteria)
    stats = assign_flags(recs, catalogue)
    print(f'read {len(recs)} condition rows from {os.path.basename(a.master)} '
          f'| GLOBAL rules available: {len(catalogue)}')
    print(f"  flags inherited from GLOBAL: {stats[cm.ORIG_GLOBAL]} "
          f"| classified automatically: {stats[cm.ORIG_AUTO]}")

    if a.dry_run:
        tables = cm.normalise(recs, cm.FONTE_LURH)
        print('  ' + ' | '.join(f'{n} {len(t)}' for n, t in zip(
            ['catálogo', 'textos', 'factos', 'licenças', 'parâmetros'], tables)))
        return 0

    if os.path.exists(a.dst):
        bak = cm.backup(a.dst)
        print(f'backup -> {os.path.basename(bak)}')

    extra = [
        ('', ''),
        ('Origem dos dados', 'Tabela de factos derivada de master_lurh.xlsx -> Conditions. '
                             'Os textos legais vêm do próprio master; os critérios vêm de '
                             'CriteriosConformidade_GLOBAL quando há correspondência.'),
        ('Por validar', f"{stats[cm.ORIG_AUTO]} linhas têm critérios classificados "
                        f"automaticamente ('{cm.ORIG_AUTO}', assinaladas a laranja no "
                        'Catalogo_Criterios). Validar antes de usar em relatórios.'),
        ('Atenção', "Em CriteriosConformidade_GLOBAL a coluna 'VLE' é uma FLAG 0/1. "
                    "Neste modelo o limite numérico é 'VLE' e a flag é 'VLE (flag)'."),
    ]
    *tables, vista = cm.write_workbook(recs, cm.FONTE_LURH, a.dst, readme_extra=extra)
    problems = cm.check_integrity(*tables)
    print(f'wrote {os.path.basename(a.dst)}: ' + ' | '.join(
        f'{n} {len(t)}' for n, t in zip(['catálogo', 'textos', 'factos', 'licenças',
                                         'parâmetros'], tables)))
    print('integrity: ' + ('OK' if not problems else '; '.join(problems)))
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
