# -*- coding: utf-8 -*-
"""CODMAXIMO — the asset key that links each licence to the ETAR registry.

The licences carry no asset code; they name the installation in free text. The
registry (TabelaCentroide_ETAR_newDep_sensity.xlsx) holds CODMAXIMO, coordinates
and every asset attribute, keyed by DESIGNACAO_GNA. We attach CODMAXIMO to each
licence with a two-stage match:

    1. NAME  — fold the name (accents/articles/"ETAR" prefix dropped, Stº->Santo)
               and match DESIGNACAO_GNA. The licence's declared identity wins, so
               a licence whose coordinates are wrong still lands on the right plant.
    2. COORD — for the residue, the nearest registry asset within 2 km (WGS84).

When a plant has both a current and an "(Antiga)" registry record, the current
one is preferred. Validated at 269/269 licences (265 name+GPS agree < 0.5 km).

REGISTRY SOURCE OF TRUTH: data/powerbi/TabelaCentroide_ETAR_newDep_sensity.xlsx
"""
import math
import os
import re
import unicodedata

import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))   # project root (this module is in src/Common_Code_PowerBI)
REGISTRY_XLSX = os.path.join(ROOT, 'data', 'powerbi', 'TabelaCentroide_ETAR_newDep_sensity.xlsx')
LICENSES_XLSX = os.path.join(ROOT, 'data', 'powerbi', 'Licenses.xlsx')

_ARTIGOS = {'DE', 'DA', 'DO', 'DAS', 'DOS', 'D'}
_EXPAND = {'STO': 'SANTO', 'STA': 'SANTA'}
_COORD_MAX_KM = 2.0


def _fold(txt):
    t = unicodedata.normalize('NFKD', str(txt)).encode('ascii', 'ignore').decode()
    toks = [_EXPAND.get(w, w) for w in re.sub(r'[^A-Z0-9 ]', ' ', t.upper()).split()
            if w not in _ARTIGOS]
    return ''.join(toks)


def _key(txt):
    """Folded key with a leading ETAR/ETA token removed (both sides carry it)."""
    k = _fold(txt)
    for pre in ('ETAR', 'ETA', 'EETAR'):
        if k.startswith(pre):
            return k[len(pre):]
    return k


def _num(v):
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return None
    try:
        return float(str(v).replace(',', '.').strip())
    except (ValueError, AttributeError):
        return None


def _haversine(la1, lo1, la2, lo2):
    r = 6371.0
    p = math.pi / 180
    a = (math.sin((la2 - la1) * p / 2) ** 2
         + math.cos(la1 * p) * math.cos(la2 * p) * math.sin((lo2 - lo1) * p / 2) ** 2)
    return 2 * r * math.asin(min(1.0, math.sqrt(a)))


def load_registry(path=REGISTRY_XLSX):
    """Return a dict with the name index and the coordinate list for matching."""
    tc = pd.read_excel(path)
    tc = tc.copy()
    tc['__key'] = tc['DESIGNACAO_GNA'].map(_key)
    tc['__la'] = tc['LATITUDE_WGS84'].map(_num)
    tc['__lo'] = tc['LONGITUDE_WGS84'].map(_num)
    tc['__antiga'] = tc['DESIGNACAO_GNA'].astype(str).str.contains('Antiga', case=False, na=False)
    by_name = {}
    for _, row in tc.iterrows():
        by_name.setdefault(row['__key'], []).append(row)
    current_by_base = {}
    for _, row in tc[~tc['__antiga']].iterrows():
        current_by_base[str(row['DESIGNACAO_GNA']).strip()] = row['CODMAXIMO']
    coords = [(r['__la'], r['__lo'], r['CODMAXIMO'], str(r['DESIGNACAO_GNA']))
              for _, r in tc.iterrows() if r['__la'] is not None and r['__lo'] is not None]
    return {'by_name': by_name, 'coords': coords, 'current_by_base': current_by_base}


def _prefer_current(codmaximo, gna_name, reg):
    if '(Antiga)' not in str(gna_name):
        return codmaximo
    base = re.sub(r'\s*\(Antiga\)', '', str(gna_name)).strip()
    return reg['current_by_base'].get(base, codmaximo)


def match_codmaximo(designacao, latitude, longitude, reg):
    """Return (codmaximo, method) for one licence, or ('', None) if unmatched."""
    la, lo = _num(latitude), _num(longitude)
    hits = reg['by_name'].get(_key(designacao), []) if _key(designacao) else []
    if hits:
        if len(hits) > 1 and la is not None:
            hits = sorted(hits, key=lambda r: _haversine(la, lo, r['__la'], r['__lo'])
                          if r['__la'] is not None else 9e9)
        row = hits[0]
        return _prefer_current(row['CODMAXIMO'], row['DESIGNACAO_GNA'], reg), 'name'
    if la is not None:
        best, best_d, best_name = None, 1e9, None
        for rla, rlo, code, name in reg['coords']:
            d = _haversine(la, lo, rla, rlo)
            if d < best_d:
                best, best_d, best_name = code, d, name
        if best is not None and best_d <= _COORD_MAX_KM:
            return _prefer_current(best, best_name, reg), 'coord'
    return '', None


def stamp_licenses_xlsx(licenses_path=LICENSES_XLSX, registry_path=REGISTRY_XLSX):
    """Add/refresh a CODMAXIMO column on Licenses.xlsx in place (openpyxl)."""
    import openpyxl
    reg = load_registry(registry_path)
    wb = openpyxl.load_workbook(licenses_path)
    ws = wb['Licenses'] if 'Licenses' in wb.sheetnames else wb.active
    hdr = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    col_des, col_lat, col_lon = hdr.get('Designação'), hdr.get('Latitude'), hdr.get('Longitude')
    col_cod = hdr.get('CODMAXIMO') or (ws.max_column + 1)
    ws.cell(1, col_cod).value = 'CODMAXIMO'
    n = matched = 0
    for r in range(2, ws.max_row + 1):
        des = ws.cell(r, col_des).value if col_des else None
        lat = ws.cell(r, col_lat).value if col_lat else None
        lon = ws.cell(r, col_lon).value if col_lon else None
        if des is None and lat is None:
            continue
        n += 1
        code, _method = match_codmaximo(des, lat, lon, reg)
        ws.cell(r, col_cod).value = code or None
        if code:
            matched += 1
    wb.save(licenses_path)
    print(f'CODMAXIMO: {matched}/{n} licences matched -> {os.path.basename(licenses_path)}')
    return matched, n


if __name__ == '__main__':
    stamp_licenses_xlsx()
