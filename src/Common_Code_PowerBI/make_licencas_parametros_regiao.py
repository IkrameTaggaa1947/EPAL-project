# -*- coding: utf-8 -*-
"""Every licensed parameter, one per row, grouped by AdVT department.

The same data as AdVT_DepDOS_VLE.xlsx, in long form instead of pivoted: rather
than one column per parameter (wide, sparse, hard to filter), one row per
(ETAR, parâmetro) with the limit next to it. That is the shape you want to
sort, filter or paste into another tool.

Four sheets, one per AdVT department (matches TabelaCentroide_ETAR's own
'Departamento DOS' values):

    DCA_S  Centro Alentejo
    DNA_S  Norte Alentejo
    DBB_S  Beira Baixa
    DBA_S  Beira Alta

Ten columns on every sheet:

    Área do CO · Município · Infraestrutura · N.º TUA · N.º LURH ·
    Data de início · Data de validade · Período de estiagem · Parâmetro · VLE

Sampling detail (local, frequência, tipo de amostragem) is deliberately NOT
here: this report answers "what limit applies", not "how is it measured".
Autocontrolo.xlsx and outputs/reports/VLE_por_ETAR.xlsx carry that.

Identity, dates, department lookup, dry-season period and the VLE cell format
are shared with make_advt_depdos_vle.py, so the two reports can never disagree.

Run:  python src/Common_Code_PowerBI/make_licencas_parametros_regiao.py
Output: outputs/reports/Licencas_Parametros_por_Regiao.xlsx
"""
import datetime
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

from make_advt_depdos_vle import (            # noqa: E402  (path set above)
    DEPARTAMENTOS,
    FEED,
    PARAM_HEADERS,
    REGISTRY_XLSX,
    find_estiagem_periods,
    licence_columns,
    read_conditions,
    read_sheet_rows,
    vle_cell,
)

ROOT = os.path.dirname(os.path.dirname(HERE))   # project root
OUT_DIR = os.path.join(ROOT, 'outputs', 'reports')
OUT = os.path.join(OUT_DIR, 'Licencas_Parametros_por_Regiao.xlsx')

COLS = ['Área do CO', 'Município', 'Infraestrutura', 'N.º TUA', 'N.º LURH',
        'Data de início', 'Data de validade', 'Período de estiagem',
        'Parâmetro', 'VLE']

WIDTHS = {'Área do CO': 22, 'Município': 16, 'Infraestrutura': 24, 'N.º TUA': 18,
          'N.º LURH': 20, 'Data de início': 12, 'Data de validade': 12,
          'Período de estiagem': 20, 'Parâmetro': 22, 'VLE': 22}


def param_label(param):
    """The parameter's display name, on one line.

    make_advt_depdos_vle.py's headers carry the line breaks the pivoted layout
    needs ('CBO5 \\n(mg/L O2)'); in a single cell they read better flattened."""
    h = PARAM_HEADERS.get(param, param)
    h = re.sub(r'\s+', ' ', h.replace('\n', ' ')).strip()
    return h.replace('NMP/ 100', 'NMP/100').replace('Escala Sorensen', 'Escala Sörensen')


def main():
    licences = {r.get('Nº TUA'): r for r in read_sheet_rows(os.path.join(FEED, 'Licenses.xlsx')) if r.get('Nº TUA')}
    if not licences:
        print('! Licenses.xlsx vazio — correr make_powerbi_all.py primeiro', file=sys.stderr)
        return 1

    registry = {r.get('CODMAXIMO'): r for r in read_sheet_rows(REGISTRY_XLSX) if r.get('CODMAXIMO')}
    estiagem_periods = find_estiagem_periods()
    cond = read_conditions()

    rows_by_dept = {d: [] for d in DEPARTAMENTOS}
    unclassified = 0

    for tua, lic in licences.items():
        cod = lic.get('CODMAXIMO')
        reg = registry.get(cod) if cod else None
        dept = str(reg.get('Departamento DOS') or '').strip() if reg else ''
        if dept not in DEPARTAMENTOS:
            unclassified += 1
            continue

        base = licence_columns(tua, lic, reg, estiagem_periods)
        params = cond.get(tua, {})
        if not params:
            # A licence with no discharge limits still gets its line, so the
            # sheet is a complete list of the department's licences.
            rows_by_dept[dept].append(dict(base, **{'Parâmetro': '', 'VLE': ''}))
            continue
        for param, by_period in params.items():
            rows_by_dept[dept].append(dict(base, **{
                'Parâmetro': param_label(param),
                'VLE': vle_cell(by_period),
            }))

    for d in rows_by_dept:
        rows_by_dept[d].sort(key=lambda r: (r['Área do CO'], r['Município'],
                                            r['Infraestrutura'], r['Parâmetro']))

    write(rows_by_dept)
    total = sum(len(v) for v in rows_by_dept.values())
    etars = sum(len({r['Infraestrutura'] for r in v}) for v in rows_by_dept.values())
    print(f'{total} linhas (ETAR x parâmetro) em {etars} infraestruturas, '
          f'{len(DEPARTAMENTOS)} departamentos')
    for d, (sheet, *_) in DEPARTAMENTOS.items():
        rows = rows_by_dept[d]
        print(f'  {sheet} ({d}): {len(rows)} linhas, '
              f'{len({r["Infraestrutura"] for r in rows})} infraestruturas')
    if unclassified:
        print(f'  {unclassified} licenças sem Departamento DOS na tabela de registo '
              f'— fora deste relatório (ver aba "Não classificado" de AdVT_DepDOS_VLE.xlsx)')
    print(f'  -> {os.path.relpath(OUT, ROOT)}')
    return 0


def write(rows_by_dept):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    os.makedirs(OUT_DIR, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    HEAD_FILL = PatternFill('solid', fgColor='1F4E78')
    HEAD_FONT = Font(name='Arial', bold=True, color='FFFFFF', size=10)
    BODY_FONT = Font(name='Arial', size=10)
    STRIPE = PatternFill('solid', fgColor='F2F6FA')
    THIN = Side(style='thin', color='D9D9D9')
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    for dept, (sheet_name, *_) in DEPARTAMENTOS.items():
        ws = wb.create_sheet(sheet_name)

        for j, col in enumerate(COLS, start=1):
            c = ws.cell(1, j, col)
            c.font = HEAD_FONT
            c.fill = HEAD_FILL
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            c.border = BORDER
        ws.row_dimensions[1].height = 22

        for i, row in enumerate(rows_by_dept[dept], start=2):
            for j, col in enumerate(COLS, start=1):
                v = row.get(col, '')
                if col in ('Data de início', 'Data de validade') and isinstance(v, datetime.datetime):
                    v = v.date()
                c = ws.cell(i, j, v if v not in ('', None) else None)
                c.font = BODY_FONT
                c.border = BORDER
                c.alignment = Alignment(wrap_text=True, vertical='center')
                if col in ('Data de início', 'Data de validade') and v:
                    c.number_format = 'DD/MM/YYYY'
                if i % 2 == 0:
                    c.fill = STRIPE

        for j, col in enumerate(COLS, start=1):
            ws.column_dimensions[get_column_letter(j)].width = WIDTHS[col]
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = f'A1:{get_column_letter(len(COLS))}{max(ws.max_row, 1)}'

    wb.save(OUT)


if __name__ == '__main__':
    sys.exit(main())
