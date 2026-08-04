"""Rebuild criteria_table_TUA.xlsx in the canonical criteria model.

The heavy lifting (schema, IDs, normalisation, styling, integrity checks) lives
in extraction/criteria_model.py, which the LURH builder imports too — that is
what guarantees the two workbooks have identical sheets and columns and can be
appended in Power BI without any mapping.

This file only knows how to READ the original flat TUA workbook
(Textos_Fonte + Criterios_0_1, 544 rows) and turn it into canonical records.
Those 544 rows were curated by hand, so their criteria flags are carried over
as-is and marked 'Manual (validado)'.

Run:  python build_criteria_table.py [--dry-run] [--in path.xlsx] [--out path.xlsx]
      python build_criteria_table.py --verify path.xlsx   (round-trip check)
A timestamped .bak_ copy is written before overwriting.
"""
import argparse
import os
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import criteria_model as cm      # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_BOOK = os.getenv(
    'EPAL_CRITERIA_TUA',
    os.path.abspath(os.path.join(HERE, '..', '..', 'data', 'reference', 'criteria_table_TUA.xlsx')))

SRC_TEXTS, SRC_FLAGS = 'Textos_Fonte', 'Criterios_0_1'


def read_flat(path):
    """The ORIGINAL two flat sheets -> canonical records."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for s in (SRC_TEXTS, SRC_FLAGS):
        if s not in wb.sheetnames:
            raise SystemExit(f"'{s}' not found in {os.path.basename(path)} — this reader expects "
                             f'the ORIGINAL flat workbook. Sheets present: {wb.sheetnames}')
    tf = list(wb[SRC_TEXTS].iter_rows(values_only=True))
    cr = list(wb[SRC_FLAGS].iter_rows(values_only=True))
    wb.close()
    if len(tf) != len(cr):
        raise SystemExit(f'sheet length mismatch: {SRC_TEXTS}={len(tf)} {SRC_FLAGS}={len(cr)}')

    ti = {h: i for i, h in enumerate(tf[0])}
    ci = {h: i for i, h in enumerate(cr[0])}
    missing = [c for c in cm.FLAG_COLS if c not in ci]
    if missing:
        raise SystemExit(f'{SRC_FLAGS} is missing flag columns: {missing}')

    rows = []
    for n, (t, c) in enumerate(zip(tf[1:], cr[1:]), start=2):
        if (t[ti['Nº TUA']], t[ti['Parâmetro (catálogo)']]) != \
           (c[ci['Nº TUA']], c[ci['Parâmetro']]):
            raise SystemExit(f'row {n}: {SRC_TEXTS} and {SRC_FLAGS} are not aligned')
        rows.append({
            'IdLicenca': t[ti['Nº TUA']],
            'Estabelecimento': t[ti['ETAR']],
            'Empresa': t[ti['Empresa']],
            'Ficheiro fonte': t[ti['Ficheiro fonte (TUA)']],
            'Regime': 'Normal',                       # not a TUA concept
            'Parâmetro (documento)': t[ti['Parâmetro (TUA)']],
            'Parâmetro (catálogo)': t[ti['Parâmetro (catálogo)']],
            'VLE': t[ti['VLE']],
            'VLE (% mín. remoção)': None,             # not carried by the TUA table
            'Nota/Código': t[ti['Nota/Código']],
            'Interpretação → critérios': t[ti['Interpretação → critérios']],
            'Texto da avaliação de conformidade (lei)':
                t[ti['Texto da avaliação de conformidade (lei)']],
            'Legislação aplicável (texto)': t[ti['Legislação aplicável (texto)']],
            'Empresa (catálogo)': c[ci['Empresa (catálogo)']] if 'Empresa (catálogo)' in ci else None,
            'Origem dos critérios': cm.ORIG_MANUAL,
            'flags': tuple(c[ci[f]] for f in cm.FLAG_COLS),
        })
    return rows


def read_canonical(path):
    """An ALREADY-canonical workbook -> records (so the build is idempotent)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    vp = list(wb['Vista_Plana'].iter_rows(values_only=True))
    wb.close()
    ix = {h: i for i, h in enumerate(vp[0])}
    rows = []
    for r in vp[1:]:
        rec = {k: r[ix[k]] for k in
               ['IdLicenca', 'Estabelecimento', 'Empresa', 'Ficheiro fonte', 'Regime',
                'Parâmetro (documento)', 'Parâmetro (catálogo)', 'VLE', 'VLE (% mín. remoção)',
                'Nota/Código', 'Interpretação → critérios',
                'Texto da avaliação de conformidade (lei)', 'Legislação aplicável (texto)',
                'Empresa (catálogo)', 'Origem dos critérios']}
        rec['flags'] = tuple(r[ix[f]] for f in cm.FLAG_COLS)
        rows.append(rec)
    return rows


def read_source(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    names = wb.sheetnames
    wb.close()
    return read_canonical(path) if 'Vista_Plana' in names else read_flat(path)


def verify_against_flat(canonical_path, flat_path):
    """Every original row must reappear, cell for cell, in Vista_Plana."""
    orig = read_flat(flat_path)
    wb = openpyxl.load_workbook(canonical_path, read_only=True, data_only=True)
    vp = list(wb['Vista_Plana'].iter_rows(values_only=True))
    wb.close()
    ix = {h: i for i, h in enumerate(vp[0])}
    cols = ['IdLicenca', 'Estabelecimento', 'Empresa', 'Ficheiro fonte',
            'Parâmetro (documento)', 'Parâmetro (catálogo)', 'VLE', 'Nota/Código',
            'Interpretação → critérios', 'Texto da avaliação de conformidade (lei)',
            'Legislação aplicável (texto)']
    diffs = []
    for i, (o, n) in enumerate(zip(orig, vp[1:]), start=2):
        got = tuple([n[ix[c]] for c in cols] + [n[ix[f]] for f in cm.FLAG_COLS])
        want = tuple([o[c] for c in cols] + list(o['flags']))
        if got != want:
            diffs.append((i, want, got))
    return len(orig), len(vp) - 1, diffs


def main(argv):
    ap = argparse.ArgumentParser()
    ap.add_argument('--in', dest='src', default=DEFAULT_BOOK)
    ap.add_argument('--out', dest='dst', default=None)
    ap.add_argument('--dry-run', action='store_true')
    ap.add_argument('--verify', metavar='CANONICAL.xlsx',
                    help='compare a built workbook against the flat source in --in')
    a = ap.parse_args(argv)

    if a.verify:
        n_orig, n_new, diffs = verify_against_flat(a.verify, a.src)
        print(f'original {n_orig} rows | rebuilt {n_new} rows | differences {len(diffs)}')
        return 0 if not diffs and n_orig == n_new else 1

    dst = a.dst or a.src
    rows = read_source(a.src)
    print(f'read {len(rows)} rows from {os.path.basename(a.src)}')

    if a.dry_run:
        tables = cm.normalise(rows, cm.FONTE_TUA)
        print('  ' + ' | '.join(f'{n} {len(t)}' for n, t in zip(
            ['catálogo', 'textos', 'factos', 'licenças', 'parâmetros'], tables)))
        return 0

    if dst == a.src:
        bak = cm.backup(dst)
        if bak:
            print(f'backup -> {os.path.basename(bak)}')

    *tables, vista = cm.write_workbook(rows, cm.FONTE_TUA, dst)
    problems = cm.check_integrity(*tables)
    print(f'wrote {os.path.basename(dst)}: ' + ' | '.join(
        f'{n} {len(t)}' for n, t in zip(['catálogo', 'textos', 'factos', 'licenças',
                                         'parâmetros'], tables)))
    print('integrity: ' + ('OK' if not problems else '; '.join(problems)))
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
