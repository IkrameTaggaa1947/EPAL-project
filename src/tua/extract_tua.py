"""
Reliable TUA extractor.

Strategy (resilient to layout shifts, unlike the old coordinate-based code):
  * Tables are read with PyMuPDF `find_tables()` (the SILIAMB TUA PDFs are ruled).
  * Columns are located by HEADER NAME (space-insensitive), never by x-position.
  * Sections are located by their MARKER text ("EXP8.3.13", ...), so e.g. the
    real "Condições de Rejeição" is never confused with the "...no ano de arranque" table.
  * Cell text is de-wrapped so words split across narrow columns are rejoined.

Run:  python extract_tua.py [path/to/TUA.pdf]
Outputs (in the configured output dir): <name>_extracted.json + <name>_extracted.xlsx
"""
import fitz
import re
import os
import sys
import json
import shutil
import unicodedata
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

DATE = re.compile(r'\d{2}-\d{2}-\d{4}')
TCODE = re.compile(r'\bT\d{6}\b')

# 'siliamb.apambiente' (the footer URL), NOT bare 'siliamb' — real data rows in
# OBRIGAÇÕES DE COMUNICAÇÃO legitimately say 'Plataforma SILiAmb'.
BOILER =['estado: emitido', 'data de consulta', 'para realizar a valida',
          'codigo documento', 'codigo verifica', 'siliamb.apambiente', 'unico ambiental',
          'validar t', 'pag.']


# ----------------------------------------------------------------- helpers
def norm(s):
    if not s:
        return ''
    s = unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', s).strip().lower()


def _key(s):
    return norm(s).replace(' ', '')


def clean_text(s):
    return re.sub(r'\s+', ' ', (s or '')).strip()


def clean_code(s):
    return re.sub(r'\s+', '', (s or ''))


def dewrap(txt):
    """Join the physical lines of a cell. A line ending in whitespace marks a word
    boundary (join with a space); a line with no trailing space was split by the
    wrap (join with none). This signal comes straight from the PDF, so it cleanly
    separates "gradagem |mecânica" (space) from "secundári|o" (no space)."""
    if not txt:
        return ''
    out, prev_boundary = '', True
    for ln in txt.split('\n'):
        boundary = ln != ln.rstrip()          # did this physical line end at a word boundary?
        seg = ln.strip()
        if not seg:
            continue
        out = seg if not out else (out + ' ' + seg if prev_boundary else out + seg)
        prev_boundary = boundary
    return re.sub(r'\s+', ' ', out).strip()


def col_idx(header, *names):
    for i, c in enumerate(header):
        kc = _key(c)
        if kc and any(_key(n) in kc for n in names):
            return i
    return None


def col_exact(header, target):
    t = _key(target)
    for i, c in enumerate(header):
        if _key(c) == t:
            return i
    return None


def find_header(rows, *required, start=0):
    for i in range(start, len(rows)):
        if all(col_idx(rows[i], n) is not None for n in required):
            return i
    return None


def find_marker(rows, marker, start=0):
    # marker is a regex; the section chapter varies (EXP8.3.x vs EXP9.3.x).
    m = norm(marker)
    for i in range(start, len(rows)):
        if re.search(m, norm(' '.join(rows[i]))):
            return i
    return None


def first_data_row(rows, start, predicate):
    for i in range(start, len(rows)):
        if predicate(rows[i]):
            return rows[i]
    return None


def _row(dr, ci, code_keys):
    out = {}
    for k, i in ci.items():
        if i is None or i >= len(dr):
            out[k] = ''
        else:
            out[k] = clean_code(dr[i]) if k in code_keys else clean_text(dr[i])
    return out


# ----------------------------------------------------------------- sections
def dados_gerais(rows):
    labels = [('no tua', 'Nº TUA'), ('estabelecimento', 'Estabelecimento'),
              ('codigo apa', 'Código APA'), ('requerente', 'Requerente'),
              ('identificacao fiscal', 'NIF'), ('localizacao', 'Localização')]
    out = {}
    for r in rows:
        if len(r) >= 2:
            key, val = norm(r[0]), clean_text(r[1])
            for sub, label in labels:
                if sub in key and val and label not in out:
                    out[label] = val
    return out


def enquadramento(rows):
    h = find_header(rows, 'data de validade', 'sentido da decisao')
    if h is None:
        return {}
    header = rows[h]
    ci = {
        'Nº Processo': col_idx(header, 'no processo'),
        'Data de Emissão': col_idx(header, 'data de emissao'),
        'Data de Entrada em Vigor': col_idx(header, 'data de entrada em vigor'),
        'Data de Validade': col_idx(header, 'data de validade'),
        'Sentido da decisão': col_idx(header, 'sentido da decisao'),
        'Entidade Licenciadora': col_idx(header, 'entidade licenciado'),
    }
    dr = first_data_row(rows, h + 1, lambda r: bool(DATE.search(' '.join(r))))
    if dr is None:
        return {}
    out = _row(dr, ci, set())
    # Nº Processo: drop any appended TURH code, rejoin wrap fragments, then drop
    # a short trailing indicador digit (the cell layout varies between documents).
    proc = re.sub(r'\s*-\s*L\d.*$', '', out.get('Nº Processo', '')).strip()
    toks = proc.split()
    if len(toks) >= 2 and re.fullmatch(r'\d{1,2}', toks[-1]):
        toks = toks[:-1]
    m = re.search(r'PL\d+', ''.join(toks))
    if m:
        out['Nº Processo'] = m.group(0)
    # dates: drop stray wrap spaces, e.g. "19-04- 2028" -> "19-04-2028"
    for dk in ('Data de Emissão', 'Data de Entrada em Vigor', 'Data de Validade'):
        if out.get(dk):
            out[dk] = re.sub(r'\s+', '', out[dk])
    return out


ENQ_TEXT = re.compile(
    r'Data\s+de\s+Emiss\S+\s+Data\s+de\s+Entrada\s+em\s+Vigor\s+Data\s+de\s+Validade\s+'
    r'(?:L[\d.]+\.?\s*RH[\w.]*\s*V?\d*\s+)?'
    r'(\d{2}-\d{2}-\d{4})\s+(\d{2}-\d{2}-\d{4})\s+(\d{2}-\d{2}-\d{4})')


def enquadramento_from_text(full_text):
    """Fallback for the ENQUADRAMENTO layout variant that has no 'Sentido da
    decisão' column (so the header lookup fails) and whose Validade column falls
    outside the detected grid: read the three dates from the text layer, where
    they always appear right after the three date headers."""
    out = {}
    m = ENQ_TEXT.search(full_text)
    if m:
        (out['Data de Emissão'], out['Data de Entrada em Vigor'],
         out['Data de Validade']) = m.groups()
    p = re.search(r'\bPL\d{8,}\b', full_text)
    if p:
        out['Nº Processo'] = p.group(0)
    return out


def localizacao_833(rows):
    # header names vary by layout generation: 'Código TURH' (EXP8.3.3) vs
    # 'Código Utilização' (Const23.2.1) — accept either.
    h = find_header(rows, 'codigo turh', 'massa de agua', 'classificacao da massa')
    if h is None:
        h = find_header(rows, 'codigo utilizacao', 'massa de agua', 'classificacao da massa')
    if h is None:
        return {}
    header = rows[h]
    ci = {
        'Código': col_idx(header, 'codigo'),
        'Código TURH': col_idx(header, 'codigo turh', 'codigo utilizacao'),
        'Longitude': col_idx(header, 'longitude'),
        'Latitude': col_idx(header, 'latitude'),
        'Massa de Água': col_idx(header, 'massa de agua'),
        'Classificação da Massa de Água': col_idx(header, 'classificacao da massa'),
    }
    dr = first_data_row(rows, h + 1, lambda r: bool(TCODE.search(' '.join(r))))
    if dr is None:
        return {}
    return _row(dr, ci, {'Código', 'Código TURH', 'Longitude', 'Latitude'})


def caracterizacao_835(rows):
    h = find_header(rows, 'designacao', 'ano de arranque', 'caudal maximo de descarga')
    if h is None:
        return {}
    header = rows[h]
    ci = {
        'Designação': col_idx(header, 'designacao'),
        'Ano de arranque': col_idx(header, 'ano de arranque'),
        'População servida à data': col_idx(header, 'servida a data'),
        'Ano horizonte de projeto': col_idx(header, 'ano horizonte'),
        'População servida no horizonte': col_idx(header, 'servida no ano'),
        'Nível de tratamento': col_idx(header, 'nivel de tratament'),
        'Esquema de tratamento': col_idx(header, 'esquema de tratament'),
        'Caudal máximo de descarga': col_idx(header, 'caudal maximo de descarga'),
    }
    dr = first_data_row(rows, h + 1, lambda r: bool(TCODE.search(' '.join(r))))
    if dr is None:
        return {}
    nums = {'Ano de arranque', 'População servida à data',
            'Ano horizonte de projeto', 'População servida no horizonte'}
    return _row(dr, ci, nums)


def _ci(header, keys):
    ci = {}
    for k, terms in keys.items():
        ci[k] = col_exact(header, 'vle') if k == 'VLE' else col_idx(header, *terms)
    return ci


def _collect(rows, marker, stop_regex, header_required, keys):
    """Collect a parameter table that may span pages (re-mapping columns on each
    header fragment) and whose parameter names may be split across a page break."""
    m = find_marker(rows, marker)
    if m is None:
        return []
    out, ci, pending = [], None, ''
    for i in range(m + 1, len(rows)):
        r = rows[i]
        joined = ' '.join(r)
        nj = norm(joined)
        if re.search(stop_regex, nj):
            break
        if any(b in nj for b in BOILER):
            continue
        if all(col_idx(r, h) is not None for h in header_required):
            ci = _ci(r, keys)                  # (re)map columns on every header fragment
            continue
        if ci is None:
            continue
        pcol = ci.get('Parâmetro')
        param = clean_text(r[pcol]) if pcol is not None and pcol < len(r) else ''
        vcol = ci.get('VLE')
        has_vle = vcol is not None and vcol < len(r) and r[vcol].strip()
        if not TCODE.search(joined):
            if param and not has_vle:          # orphan parameter fragment (page-break wrap)
                pending = (pending + ' ' + param).strip()
            continue
        rec = _row(r, ci, {'Código', 'Código TURH'})
        if 'VLE' in keys and (not rec.get('Parâmetro', '').strip()
                              or _REF_CELL.fullmatch(clean_code(rec.get('VLE', '')))):
            rec = _repair_shifted_row(r, rec)
        if pending:
            rec['Parâmetro'] = (pending + ' ' + rec.get('Parâmetro', '')).strip()
            pending = ''
        out.append(rec)
    return out


CONDITION_KEYS = {
    'Código': ('codigo',), 'Código TURH': ('codigo turh',), 'Parâmetro': ('parametro',),
    'VLE (% mín. redução)': ('reducao',), 'VLE': (),
    'Carga máx. admissível (kg/dia)': ('carga max',),
    'Legislação aplicável': ('legislacao aplicavel',),
    'Avaliação da conformidade': ('avaliacao da conformidade',),
    'Observações': ('observacoes',),
}
AUTOCONTROLO_KEYS = {
    'Código': ('codigo',), 'Código TURH': ('codigo turh',),
    'Local de amostragem': ('local de amostragem',), 'Parâmetro': ('parametro',),
    'Frequência de amostragem': ('frequencia de amostragem',),
    'Tipo de amostragem': ('tipo de amostragem',), 'Observações': ('observacoes',),
}


_REF_CELL = re.compile(r'^\([a-z0-9]\)$|^T\d{6}$')
_NUM_CELL = re.compile(r'^\d+(?:[.,]\d+)?(?:\s*(?:-|a)\s*\d+(?:[.,]\d+)?)?$')


def _repair_shifted_row(r, rec):
    """2025-layout continuation pages drop the table header and shift the grid
    columns, so the page-1 column map reads the wrong cells (Parâmetro empty,
    VLE showing '(a)' or a T-code). Rebuild the record from cell CONTENT."""
    cells = [clean_text(c) for c in r if clean_text(c)]
    if len(cells) < 4 or not re.fullmatch(r'T\d{6}', clean_code(cells[0])):
        return rec
    param = vle = turh = ''
    refs = []
    for c in cells[1:]:
        cc = clean_code(c)
        if re.match(r'^L\d{4,}', cc) and not turh:
            turh = cc
        elif _REF_CELL.fullmatch(cc):
            refs.append(cc)
        elif _NUM_CELL.fullmatch(c.strip()) and not vle:
            vle = c.strip()
        elif re.search(r'[A-Za-zÀ-ÿ]{3,}', c) and not param:
            param = c
    if not param and not vle:
        return rec
    fixed = {k: '' for k in rec}
    fixed.update({'Código': clean_code(cells[0]), 'Código TURH': turh,
                  'Parâmetro': param, 'VLE': vle})
    if refs:
        fixed['Legislação aplicável'] = refs[0]
    if len(refs) > 1:
        fixed['Avaliação da conformidade'] = refs[1]
    return fixed


def conditions_8313(rows):
    return _collect(rows, r'exp[89]\.3\.13', r'exp[89]\.3\.1[4-9]', ['parametro', 'avaliacao'], CONDITION_KEYS)


def autocontrolo_8316(rows):
    return _collect(rows, r'exp[89]\.3\.16', r'exp[89]\.3\.1[7-9]|exp[89]\.3\.2', ['parametro', 'frequencia'], AUTOCONTROLO_KEYS)


def _dedupe_trailing(s):
    """Drop a duplicated trailing clause (a tall-cell get_textbox artifact)."""
    w = s.split()
    n = len(w)
    for k in range(min(30, n // 2), 6, -1):
        tail = w[n - k:]
        for j in range(0, n - k):
            if w[j:j + k] == tail:               # the last k words recur earlier
                return ' '.join(w[:n - k]).rstrip(' ,;.') + '.'
    return s


def legend(rows, marker, next_markers):
    """code -> full text. Handles BOTH schemes: '(a)/(b)' footnote letters
    embedded in the text, and per-row legend T-codes used as the key."""
    m = find_marker(rows, marker)
    if m is None:
        return {}
    tcode_map, parts = {}, []
    for i in range(m + 1, len(rows)):
        cells = [c for c in rows[i] if c]
        if not cells:
            continue
        full = norm(' '.join(cells))
        if any(re.search(nm, full) for nm in next_markers):
            break
        if 'codigo turh' in full or any(b in full for b in BOILER):
            continue
        text = max(cells, key=len)
        if norm(text) in ('avaliacao da conformidade', 'legislacao aplicavel', 'condicao'):
            continue
        text = re.sub(r'\bT\d{6}\b|L\d{6}\.\d{4}\.\s*\w+\.\w+', '', text).strip()
        codes = re.findall(r'\bT\d{6}\b', ' '.join(cells))
        key = codes[0] if codes else (parts[-1][0] if parts else None)  # continuation -> previous T-code
        if key:
            tcode_map[key] = (tcode_map.get(key, '') + ' ' + text).strip()
        parts.append((key, text))
    # scheme A: split the joined text on (a)/(b)/(1)/(2)...
    joined = re.sub(r'\s+', ' ', ' '.join(t for _, t in parts))
    letter_map, cur = {}, None
    for p in re.split(r'(\([a-z0-9]\))', joined):
        mm = re.fullmatch(r'\(([a-z0-9])\)', p.strip()) if p.strip() else None
        if mm:
            cur = mm.group(1)
        elif cur is not None:
            letter_map[cur] = (letter_map.get(cur, '') + ' ' + p).strip()
    out = {k: _dedupe_trailing(re.sub(r'\s+', ' ', v).strip()) for k, v in tcode_map.items()}
    for k, v in letter_map.items():
        out[k] = _dedupe_trailing(re.sub(r'\s+', ' ', v).strip())
    return out


def resolve(cell, legend_map):
    """Resolve a Legislação/Avaliação cell to full text. The cell may reference
    footnote letters '(a)' or legend T-codes 'T000007' (varies by document)."""
    keys = re.findall(r'T\d{6}', cell or '') + re.findall(r'\(([a-z0-9])\)', cell or '')
    out = []
    for k in keys:
        if k in legend_map:
            out.append(legend_map[k])
        elif len(legend_map) == 1:
            # the row references a code the TUA never defines (seen in Alverca da
            # Beira: rows cite T000010, section defines only T000011) — when the
            # section holds exactly one text, that is the intended one.
            out.append(next(iter(legend_map.values())))
        else:
            out.append(k)
    return ' | '.join(out)


# ------------------------------------------------ generic per-section extraction
# The curated functions above pull the fields Power BI needs. This block instead
# dumps EVERY section of the certificate generically, one table per section, so
# nothing is dropped — the "all sections" view.
SECTION_MARK = re.compile(r'exp\s*\d+\.(\d+\.\d+)\s*[-–]\s*(.*)', re.S)
# 'OBRIGAÇÕES DE COMUNICAÇÃO / OCom1 - ...' has no EXP marker; detect it by the
# 'OComN -' code so it becomes its OWN section instead of gluing onto 3.21.
# (Keying on 'OComN' — not the chapter name — keeps the page-1 table of contents,
# which lists 'OBRIGAÇÕES DE COMUNICAÇÃO' alone, from creating a bogus section.)
OCOM_MARK = re.compile(r'\bocom(\d+)\s*[-–]')

# Excel-safe, human sheet names per section code (fallback: "Sec <code>").
SECTION_NAMES = {
    '3.3': 'Localizacao', '3.5': 'Caracterizacao Geral',
    '3.7': 'Rejeicao Aguas Residuais', '3.8': 'Afluente Bruto',
    '3.11': 'Origem Aguas Residuais', '3.13': 'Condicoes Rejeicao',
    '3.14': 'Legislacao', '3.15': 'Avaliacao Conformidade',
    '3.16': 'Autocontrolo', '3.19': 'Condicoes Gerais',
    '3.20': 'Condicoes Especificas', '3.21': 'Outras Condicoes',
    'OCom1': 'Obrigacoes Comunicacao',
}


def _clean_title(text):
    line = text.strip().splitlines()[0] if text and text.strip() else ''
    line = re.split(r'\bEXP\s*\d', line)[0]          # cut a merged following chapter header
    return re.sub(r'\s+', ' ', line).strip()


def _row_y(trow, table):
    ys = [c[1] for c in trow.cells if c]
    return min(ys) if ys else table.bbox[1]


def _read_tables(doc):
    """Single pass over the ruled tables → (flat_rows, raw_sections).

    Each cell is read once with page.get_textbox — deliberately, because it keeps
    the trailing-space signal at wrapped-line boundaries that dewrap() needs; the
    ~27x faster Table.extract() drops those spaces ("ETAR Foios" -> "ETARFoios").
    That single read feeds BOTH outputs, so the expensive work is not doubled:
      * flat_rows    — every row in page/table order (what the curated field
                       extractors below consume; identical to the old all_rows).
      * raw_sections — [code, title, [rows]] split by EXP marker. Markers come
                       from the text layer (per LINE — a block can hold several),
                       and markers and individual ROWS are interleaved by their
                       y-position, so several sections rendered as ONE physical
                       table still separate, and 3.19 (whose marker is not a table
                       row) is not merged into 3.16.
    ponytail: find_tables + per-cell get_textbox is ~36s/8pp — the single pass
    keeps it to one; a faster path must still preserve the wrap-space signal."""
    flat, out, cur = [], [], None
    for page in doc:
        events = []                                  # (y, order, kind, payload), top-to-bottom
        for block in page.get_text('dict')['blocks']:
            for line in block.get('lines', []):
                text = ''.join(sp['text'] for sp in line['spans'])
                m = SECTION_MARK.search(norm(text))
                if m:
                    mm = re.search(r'EXP\s*\d+\.' + re.escape(m.group(1)) + r'\s*[-–]\s*(.*)', text)
                    events.append((line['bbox'][1], 0, 'mark',
                                   (m.group(1), _clean_title(mm.group(1) if mm else ''))))
                    continue
                mo = OCOM_MARK.search(norm(text))
                if mo:
                    mt = re.search(r'OCom\d+\s*[-–]\s*(.*)', text)
                    events.append((line['bbox'][1], 0, 'mark',
                                   (f'OCom{mo.group(1)}',
                                    _clean_title(mt.group(1) if mt else 'Comunicações'))))
        for t in page.find_tables().tables:
            for trow in t.rows:
                cells = [dewrap(page.get_textbox(b)) if b else '' for b in trow.cells]
                flat.append(cells)                   # page/table order (old all_rows)
                events.append((_row_y(trow, t), 1, 'row', cells))
        # order key breaks y-ties so a marker sorts before its own table row.
        for _, _, kind, payload in sorted(events, key=lambda e: (e[0], e[1])):
            if kind == 'mark':
                code, title = payload
                if cur and cur[0] == code:           # marker seen as both a line and a row
                    continue
                cur = [code, title, []]
                out.append(cur)
            elif cur is not None:
                cur[2].append(payload)
    return flat, out


def _section_table(srows):
    """(columns, [row dicts]) from a section's raw grid, cleaned for reading:

    * the header is the first row that labels the 'Código' column (every SILIAMB
      section table starts there), so a stray banner row above it is skipped;
    * boilerplate and repeated header rows are dropped;
    * all-empty columns are dropped — find_tables emits phantom columns between the
      real ones (a 22-column grid is really ~10 columns), and those are pure noise.
    Both the columns list and the row dicts are trimmed, so the per-file sheet and
    the consolidated workbook are equally clean."""
    clean = []
    for r in srows:
        nj = norm(' '.join(c for c in r if c))
        if (not nj or SECTION_MARK.search(nj) or OCOM_MARK.search(nj)
                or 'obrigacoes de comunicacao' in nj or any(b in nj for b in BOILER)):
            continue
        clean.append(r)
    if not clean:
        return [], []
    hi = next((i for i, r in enumerate(clean) if any(norm(c) == 'codigo' for c in r)), 0)
    head = clean[hi]
    width = max(len(r) for r in clean[hi:])
    header = [(head[i].strip() if i < len(head) else '') or f'Coluna {i + 1}' for i in range(width)]
    seen = {}
    for i, c in enumerate(header):                    # de-duplicate column names for Excel
        seen[c] = seen.get(c, 0) + 1
        if seen[c] > 1:
            header[i] = f'{c} ({seen[c]})'
    data = []
    for r in clean[hi + 1:]:
        if r == head:                                 # header repeated across a page break
            continue
        r = list(r) + [''] * (width - len(r))
        data.append({header[i]: clean_text(r[i]) for i in range(width)})

    def keep_col(c):
        filled = sum(1 for row in data if (row.get(c) or '').strip())
        # unnamed phantom columns (find_tables emits them between real ones) are kept
        # only when they clearly hold repeated real data, not a stray wrap fragment.
        return filled >= (2 if c.startswith('Coluna ') else 1)
    keep = [c for c in header if keep_col(c)]
    return keep, [{c: row.get(c, '') for c in keep} for row in data]


def _build_sections(raw_sections):
    """[{code, title, sheet, columns, rows}] — the generic per-section dump, from
    the raw [code, title, grid] produced by _read_tables."""
    result = []
    for code, title, srows in raw_sections:
        cols, rows = _section_table(srows)
        if not rows:
            continue
        result.append({'code': code, 'title': title,
                       'sheet': SECTION_NAMES.get(code, f'Sec {code}'),
                       'columns': cols, 'rows': rows})
    return result


# 3.13 (conditions) and 3.16 (autocontrolo) are complex multi-column tables that
# find_tables mis-grids (phantom columns, split cells). We already parse them
# cleanly by header name, so those parsed rows drive the section sheet instead.
_CURATED_COLS = {
    '3.13': ['Código', 'Código TURH', 'Parâmetro', 'VLE (% mín. redução)', 'VLE',
             'Carga máx. admissível (kg/dia)', 'Legislação aplicável',
             'Avaliação da conformidade', 'Observações'],
    '3.16': ['Código', 'Código TURH', 'Local de amostragem', 'Parâmetro',
             'Frequência de amostragem', 'Tipo de amostragem', 'Observações'],
}


def _curated_override(sections, curated):
    """Replace the generic dump for {code: clean row list} with only the non-empty
    curated columns, so the parameter-heavy sections read cleanly."""
    for s in sections:
        items = curated.get(s['code'])
        if items:
            keep = [c for c in _CURATED_COLS[s['code']]
                    if any((it.get(c) or '').strip() for it in items)]
            s['columns'] = keep
            s['rows'] = [{c: it.get(c, '') for c in keep} for it in items]
    return sections


# ----------------------------------------------------------------- top level
def extract(pdf_path):
    doc = fitz.open(pdf_path)
    rows, raw_sections = _read_tables(doc)            # one find_tables/get_textbox pass
    # The sampling-type footnote after the 3.16 table ("Amostragem composta
    # recolhida durante um período de 24 horas: (i)...") is FREE TEXT, not a table
    # row, so the grid pass never sees it — grab it from the text layer instead.
    full_text = '\n'.join(page.get_text() for page in doc)
    doc.close()
    mnote = re.search(r'Amostragem\s+composta\s+recolhida.*?(?=\n\s*EXP|\Z)',
                      full_text, re.S | re.I)
    nota = clean_text(mnote.group(0)) if mnote else ''
    legis = legend(rows, r'exp[89]\.3\.14', [r'exp[89]\.3\.15', 'avaliacao de conformidade'])
    aval = legend(rows, r'exp[89]\.3\.15', [r'exp[89]\.3\.16', 'programa de autocontrolo'])

    conds = conditions_8313(rows)
    for c in conds:
        c['Legislação aplicável (texto)'] = resolve(c.get('Legislação aplicável'), legis)
        c['Avaliação da conformidade (texto)'] = resolve(c.get('Avaliação da conformidade'), aval)

    data = {
        'file_name': os.path.basename(pdf_path),
        'dados_gerais': dados_gerais(rows),
        'enquadramento': enquadramento(rows),
        'localizacao': localizacao_833(rows),
        'caracterizacao': caracterizacao_835(rows),
        'condicoes_rejeicao': conds,
        'legislacao': legis,
        'avaliacao_conformidade': aval,
        'autocontrolo': autocontrolo_8316(rows),
        'autocontrolo_nota': nota,
    }
    if not (data['enquadramento'].get('Data de Entrada em Vigor')
            and data['enquadramento'].get('Data de Validade')):
        for k, v in enquadramento_from_text(full_text).items():
            if not data['enquadramento'].get(k):
                data['enquadramento'][k] = v
    data['sections'] = _curated_override(_build_sections(raw_sections),
                                         {'3.13': conds, '3.16': data['autocontrolo']})
    data['_warnings'] = _warnings(data)
    return data


def _warnings(d):
    w = []
    if not d['dados_gerais'].get('Código APA'):
        w.append('Código APA not found')
    if not d['enquadramento'].get('Data de Validade'):
        w.append('Data de Validade not found')
    if not d['condicoes_rejeicao']:
        w.append('No EXP8.3.13 discharge conditions found')
    if not d['autocontrolo']:
        w.append('No EXP8.3.16 autocontrolo rows found')
    # A parameter holding 2+ unit markers means two parameters merged (page-break
    # fragmentation) — flag it for review rather than letting it pass silently.
    unit = re.compile(r'\((?:mg/L|Escala|ºC|m3)')
    for label, items in (('condition', d['condicoes_rejeicao']), ('autocontrolo', d['autocontrolo'])):
        garbled = [it.get('Parâmetro', '') for it in items if len(unit.findall(it.get('Parâmetro', ''))) > 1]
        if garbled:
            w.append(f"{label} parameter(s) look garbled (page-break fragmentation): {garbled}")
    return w


def _status(validade):
    try:
        d = datetime.strptime(validade, '%d-%m-%Y').date()
        return 'Em vigor' if d >= datetime.now().date() else 'Caducada'
    except Exception:
        return ''


def iso_date(s):
    """'DD-MM-YYYY' -> 'YYYY-MM-DD' (a real ISO date Power BI parses without a locale)."""
    m = re.match(r'(\d{2})-(\d{2})-(\d{4})', s or '')
    return f"{m.group(3)}-{m.group(2)}-{m.group(1)}" if m else ''


def param_key(name):
    """Normalize a discharge parameter to a short, join-friendly key (CBO5, CQO, SST,
    pH, O&G, N, P). This is the column lab results join to — the licence spells the
    parameter out ('Carência Química de Oxigénio (mg/L O2)'), the lab uses 'CQO'."""
    n = norm(name)
    if 'bioqu' in n:                       return 'CBO5'   # before 'quimica' (bioQUIMICA)
    if 'quimica' in n or n == 'cqo':       return 'CQO'
    if 'solidos' in n or 'suspenso' in n:  return 'SST'
    if 'oleos' in n or 'gordura' in n:     return 'O&G'
    if 'azoto' in n:                       return 'N'
    if 'fosforo' in n:                     return 'P'
    if n.startswith('ph') or 'sorensen' in n:  return 'pH'
    return name.strip()


def vle_bounds(vle):
    """Parse a VLE cell to numeric (min, max). A single value is an upper limit
    (min=None); a range 'x a y' / 'x-y' (pH) gives both. Lets Power BI compare a
    measured value to a number instead of parsing text."""
    if not vle:
        return (None, None)
    s = str(vle).replace(',', '.')
    r = re.search(r'(-?\d+(?:\.\d+)?)\s*(?:a|-|–|ate|até)\s*(-?\d+(?:\.\d+)?)', s, re.I)
    if r:
        return (float(r.group(1)), float(r.group(2)))
    m = re.search(r'-?\d+(?:\.\d+)?', s)
    return (None, float(m.group(0))) if m else (None, None)


def tua_key(d):
    """Stable key linking every sheet to its certificate (Nº TUA, file_name fallback)."""
    return d['dados_gerais'].get('Nº TUA', '') or d['file_name']


def license_row(d):
    """The single-value certificate fields — one row per certificate (a dimension)."""
    g, e = d['dados_gerais'], d['enquadramento']
    loc, car = d['localizacao'], d['caracterizacao']
    return {
        'file_name': d['file_name'], 'TUA/LURH': 'TUA',
        'Nº TUA': g.get('Nº TUA', ''), 'Nº Processo': e.get('Nº Processo', ''),
        'Estabelecimento': g.get('Estabelecimento', ''), 'Código APA': g.get('Código APA', ''),
        'Data de Emissão': e.get('Data de Emissão', ''),
        'Data de Entrada em Vigor': e.get('Data de Entrada em Vigor', ''),
        'Data de Validade': e.get('Data de Validade', ''),
        'Validade (ISO)': iso_date(e.get('Data de Validade', '')),
        'Em Vigor/Caducada': _status(e.get('Data de Validade', '')),
        'Sentido da decisão': e.get('Sentido da decisão', ''),
        'Entidade Licenciadora': e.get('Entidade Licenciadora', ''),
        'Código TURH': loc.get('Código TURH', ''),
        'Longitude': loc.get('Longitude', ''), 'Latitude': loc.get('Latitude', ''),
        'Massa de Água': loc.get('Massa de Água', ''),
        'Classificação da Massa de Água': loc.get('Classificação da Massa de Água', ''),
        'Designação': car.get('Designação', ''),
        'Ano de arranque': car.get('Ano de arranque', ''),
        'População servida à data': car.get('População servida à data', ''),
        'Ano horizonte de projeto': car.get('Ano horizonte de projeto', ''),
        'População servida no horizonte': car.get('População servida no horizonte', ''),
        'Nível de tratamento': car.get('Nível de tratamento', ''),
        'Esquema de tratamento': car.get('Esquema de tratamento', ''),
        'Caudal máximo de descarga': car.get('Caudal máximo de descarga', ''),
    }


def condition_rows(d):
    """One row per EXP_.3.13 discharge parameter, linked to the certificate by Nº TUA."""
    tua = tua_key(d)
    saida = {norm(a['Parâmetro']): a for a in d['autocontrolo']
             if norm(a.get('Local de amostragem', '')) == 'saida'}
    out = []
    for c in d['condicoes_rejeicao']:
        ac = saida.get(norm(c.get('Parâmetro', '')), {})
        pk = param_key(c.get('Parâmetro', ''))
        lo, hi = vle_bounds(c.get('VLE', ''))
        out.append({
            'Nº TUA': tua,
            'Parâmetro': c.get('Parâmetro', ''),
            'ParamKey': pk,
            'Chave': f"{tua}|{pk}",              # composite key lab results join on
            'VLE (% mín. redução)': c.get('VLE (% mín. redução)', ''),
            'VLE': c.get('VLE', ''),
            'VLE mín': lo, 'VLE máx': hi,         # numeric bounds for a clean PBI compare
            'Carga máx. admissível (kg/dia)': c.get('Carga máx. admissível (kg/dia)', ''),
            'Legislação aplicável': c.get('Legislação aplicável (texto)', ''),
            'Avaliação da Conformidade Legal': c.get('Avaliação da conformidade (texto)', ''),
            'Observações': c.get('Observações', ''),
            'Frequência de amostragem': ac.get('Frequência de amostragem', ''),
            'Tipo de amostragem': ac.get('Tipo de amostragem', ''),
        })
    return out


def autocontrolo_rows(d):
    """Every EXP_.3.16 self-monitoring row (Saída AND Entrada, incl. monitor-only
    parameters like Azoto/Fósforo/Caudal), linked by Nº TUA. Rows whose sampling
    type carries a footnote mark like 'Composta (i)' get the footnote's full text."""
    tua, nota = tua_key(d), d.get('autocontrolo_nota', '')
    return [{'Nº TUA': tua, 'Código': a.get('Código', ''),
             'Local de amostragem': a.get('Local de amostragem', ''),
             'Parâmetro': a.get('Parâmetro', ''),
             'Frequência de amostragem': a.get('Frequência de amostragem', ''),
             'Tipo de amostragem': a.get('Tipo de amostragem', ''),
             'Nota tipo de amostragem':
                 nota if nota and re.search(r'\([ivx]+\)', a.get('Tipo de amostragem', '')) else '',
             'Observações': a.get('Observações', '')}
            for a in d['autocontrolo']]


def legend_rows(d, section, col):
    """Legend code -> full text, deduped by text (prefer the short footnote key)."""
    tua, seen, out = tua_key(d), {}, []
    for k, v in sorted(d[section].items(), key=lambda kv: len(kv[0])):
        if v not in seen:
            seen[v] = k
            out.append({'Nº TUA': tua, 'Código': k, col: v})
    return out


# Normalized master workbook — one sheet per grain, all linked on Nº TUA.
MASTER_SHEETS = ['Licenses', 'Conditions', 'Autocontrolo', 'Legislacao', 'Avaliacao']


def build_sheets(d):
    return {
        'Licenses': [license_row(d)],
        'Conditions': condition_rows(d),
        'Autocontrolo': autocontrolo_rows(d),
        'Legislacao': legend_rows(d, 'legislacao', 'Legislação aplicável'),
        'Avaliacao': legend_rows(d, 'avaliacao_conformidade', 'Avaliação da conformidade'),
    }


# --------------------------------------------------- per-file workbook (human view)
# The per-certificate '*_extracted.xlsx' is what opens in Excel for the operator, so
# it is built for READING, not for Power BI: a summary identity card + clean tables
# with only the columns a person checks against the PDF (internal codes dropped).
_HDR_FILL = PatternFill('solid', fgColor='1F4E78')     # table header: dark blue
_HDR_FONT = Font(bold=True, color='FFFFFF')
_SEC_FILL = PatternFill('solid', fgColor='D9E1F2')     # summary section band: light blue
_SEC_FONT = Font(bold=True, color='1F4E78')
_TITLE_FONT = Font(bold=True, size=14, color='1F4E78')
_WRAP = Alignment(wrap_text=True, vertical='top')
_TOP = Alignment(vertical='top')


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _table(ws, headers, rows, widths, wrap_cols=(), title=None):
    """A formatted table: an optional title banner, a styled frozen header row,
    column widths, and wrapped long-text cells."""
    hrow = 1
    if title:
        ws.append([title])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(headers)))
        ws.cell(1, 1).font = _TITLE_FONT
        ws.cell(1, 1).alignment = Alignment(vertical='center')
        ws.row_dimensions[1].height = 22
        hrow = 2
    ws.append(headers)
    for c in ws[hrow]:
        c.fill, c.font = _HDR_FILL, _HDR_FONT
        c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[hrow].height = 28
    for r in rows:
        ws.append([v if v not in (None, '') else '' for v in r])
    _widths(ws, widths)
    wrap = set(wrap_cols)
    for row in ws.iter_rows(min_row=hrow + 1):
        for cell in row:
            cell.alignment = _WRAP if cell.column in wrap else _TOP
    ws.freeze_panes = ws.cell(hrow + 1, 1).coordinate


def _summary(ws, d):
    """The first sheet: an identity card mirroring the LICENCE'S OWN ORDER — one
    block per PDF chapter/table (Dados Gerais, Enquadramento, Localização,
    Exploração), each field in the order it appears in the document. Nothing from
    two different PDF tables is mixed into one block. 'Estado' (computed) is the
    only extra, placed right after Data de Validade."""
    g, e = d['dados_gerais'], d['enquadramento']
    estado = _status(e.get('Data de Validade', ''))
    enq = []
    for k, v in e.items():
        enq.append((k, v))
        if k == 'Data de Validade':
            enq.append(('Estado', estado))
    blocks = [(title, kvs) for title, kvs in (
        ('Dados gerais', list(g.items())),
        ('Enquadramento', enq),
        ('Localização', list(d['localizacao'].items())),
        ('Exploração — caracterização', list(d['caracterizacao'].items())),
    ) if kvs]
    ws['A1'] = g.get('Estabelecimento') or d['file_name']
    ws['A1'].font = _TITLE_FONT
    ws['A2'] = f"Título Único Ambiental — {estado}" if estado else 'Título Único Ambiental'
    ws['A2'].font = Font(italic=True, color='595959')
    ws.merge_cells('A1:B1')
    ws.merge_cells('A2:B2')
    r = 4
    for title, kvs in blocks:
        ws.cell(r, 1, title).font = _SEC_FONT
        ws.cell(r, 1).fill = ws.cell(r, 2).fill = _SEC_FILL
        r += 1
        for label, val in kvs:
            ws.cell(r, 1, label).font = Font(bold=True)
            ws.cell(r, 2, val if val not in (None, '') else '—').alignment = _WRAP
            r += 1
        r += 1
    for msg in (d.get('_warnings') or []):
        ws.cell(r, 1, '⚠ Verificar').font = Font(bold=True, color='C00000')
        ws.cell(r, 2, msg).alignment = _WRAP
        r += 1
    _widths(ws, [32, 72])


_WRAP_HINTS = ('condi', 'legisla', 'avalia', 'observ', 'parametro', 'designa',
               'esquema', 'nivel de tratam', 'massa de agua', 'origem')


def _wrap_cols(cols):
    """1-indexed columns to wrap: the free-text ones, by column name."""
    return {i for i, c in enumerate(cols, 1) if any(k in norm(c) for k in _WRAP_HINTS)}


def _sheet_widths(cols, rows, wrap):
    w = []
    for i, c in enumerate(cols, 1):
        longest = max([len(str(c))] + [len(str(r.get(c, ''))) for r in rows[:200]])
        cap = 80 if i in wrap else 34
        w.append(min(cap, max(10, longest + 2)))
    return w


def _unique_sheet(name, used):
    """Excel-safe (<=31 chars, no : \\ / ? * [ ]) and unique among `used`."""
    name = re.sub(r'[:\\/?*\[\]]', ' ', name).strip()[:31] or 'Sec'
    base, n = name, 2
    while name.lower() in used:
        tag = f' {n}'
        name = base[:31 - len(tag)] + tag
        n += 1
    used.add(name.lower())
    return name


def _write_sections(wb, sections, nota=''):
    """One sheet per extracted section (clean, titled), tab named '<Title> (<code>)'.
    The 3.16 sheet gets the sampling-type footnote appended under the table."""
    used = {'resumo'}
    for s in sections:
        ws = wb.create_sheet(_unique_sheet(f"{s['sheet']} ({s['code']})", used))
        cols, rows = s['columns'], s['rows']
        wrap = _wrap_cols(cols)
        _table(ws, cols, [[r.get(c, '') for c in cols] for r in rows],
               _sheet_widths(cols, rows, wrap), wrap_cols=tuple(wrap),
               title=f"{s['code']} — {s['title']}")
        if s['code'] == '3.16' and nota:
            r = ws.max_row + 2
            ws.cell(r, 1, 'Nota').font = Font(bold=True)
            c = ws.cell(r, 2, nota)
            c.alignment = _WRAP
            if len(cols) > 2:
                ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=len(cols))
            ws.row_dimensions[r].height = 60


def write_outputs(d, out_dir):
    os.makedirs(out_dir, exist_ok=True)
    base = os.path.splitext(d['file_name'])[0]
    json_path = os.path.join(out_dir, base + '_extracted.json')
    xlsx_path = os.path.join(out_dir, base + '_extracted.xlsx')

    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(d, f, ensure_ascii=False, indent=2)

    wb = Workbook()
    _summary(wb.active, d)
    wb.active.title = 'Resumo'                      # identity card
    _write_sections(wb, d.get('sections', []),      # one sheet per section (all sections)
                    d.get('autocontrolo_nota', ''))
    wb.save(xlsx_path)
    return json_path, xlsx_path


def _style_sheet(ws, df):
    """Bold frozen header + sensible widths (and wrap on long-text columns) for a
    written data sheet — so the master workbooks read cleanly, not as raw dumps."""
    for c in ws[1]:
        c.fill, c.font = _HDR_FILL, _HDR_FONT
        c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 26
    wrap = {i for i, col in enumerate(df.columns, 1) if any(k in norm(str(col)) for k in _WRAP_HINTS)}
    for i, col in enumerate(df.columns, 1):
        longest = max([len(str(col))] + [len(str(v)) for v in df[col].head(300)])
        ws.column_dimensions[get_column_letter(i)].width = min(80 if i in wrap else 40,
                                                               max(10, longest + 2))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.column in wrap:
                cell.alignment = _WRAP


def _upsert_sheets(new_sheets, master_path, order=None, key='Nº TUA'):
    """Upsert {sheet: list-of-dicts} into an .xlsx by `key`: this certificate's
    existing rows are REPLACED (not duplicated) in every sheet. Robust to the file
    being open in Excel — never clobber an unreadable master, and divert a blocked
    write to '<name>.PENDING.xlsx' so nothing is lost.
    Returns (path_written, {sheet: nrows}, deferred)."""
    existing, read_failed = {}, False
    if os.path.isfile(master_path):
        try:
            existing = pd.read_excel(master_path, sheet_name=None)   # all sheets
        except Exception:
            read_failed = True          # locked/unreadable — do not clobber it

    names = list(dict.fromkeys((order or list(new_sheets)) + list(existing)))
    drop = {str(r[key]) for rows in new_sheets.values() for r in rows if key in r}
    result = {}
    for name in names:
        new_df = pd.DataFrame(new_sheets.get(name, []))
        old = existing.get(name, pd.DataFrame())
        if not old.empty and key in old.columns and drop:
            old = old[~old[key].astype(str).isin(drop)]   # drop this cert's old rows
        result[name] = pd.concat([old, new_df], ignore_index=True) if len(old) or len(new_df) else new_df

    pending = os.path.splitext(master_path)[0] + '.PENDING.xlsx'
    # If we couldn't read an existing master, writing would drop the other
    # certificates — divert straight to the side file instead.
    target, deferred = (pending, True) if read_failed else (master_path, False)
    # Back up the current master before overwriting it, so a bad run is recoverable.
    if not deferred and os.path.isfile(master_path):
        try:
            shutil.copy2(master_path, os.path.splitext(master_path)[0] + '.bak.xlsx')
        except Exception:
            pass

    def _write(path):
        with pd.ExcelWriter(path, engine='openpyxl') as xl:
            for name in names:
                df = result[name] if not result[name].empty else pd.DataFrame({'(sem dados)': []})
                sheet = name[:31]
                df.to_excel(xl, sheet_name=sheet, index=False)
                _style_sheet(xl.sheets[sheet], df)
    try:
        _write(target)
    except PermissionError:
        target, deferred = pending, True
        _write(target)
    return target, {n: len(result[n]) for n in names}, deferred


def update_master(d, master_path):
    """Upsert this certificate into the master workbook: the normalized Power BI
    sheets (Licenses / Conditions / Autocontrolo / Legislacao / Avaliacao) PLUS the
    same per-section sheets as each licence's own workbook — 'Localizacao (3.3)'
    ... 'Obrigacoes Comunicacao (OCom1)' — every certificate stacked, all linked
    on Nº TUA. ('Resumo' has no master twin: 'Licenses' is its one-row-per-cert
    equivalent.)"""
    sheets = dict(build_sheets(d))
    sec = section_sheets(d)
    sheets.update(sec)
    return _upsert_sheets(sheets, master_path, order=MASTER_SHEETS + list(sec))


def section_sheets(d):
    """{sheet_name: [row dicts with Nº TUA]} — one sheet per extracted section,
    named exactly like the per-file workbook tabs ('<Title> (<code>)')."""
    tua = tua_key(d)
    out = {}
    for s in d.get('sections', []):
        name = f"{s['sheet']} ({s['code']})"
        out.setdefault(name, []).extend(dict({'Nº TUA': tua}, **r) for r in s['rows'])
    return out


def update_sections_master(d, path):
    """Upsert every section of this certificate into the consolidated all-sections
    workbook (one sheet per section, all certificates, keyed on Nº TUA)."""
    return _upsert_sheets(section_sheets(d), path)


def export_master_csv(master_path, csv_dir=None):
    """Write each master sheet to a UTF-8 CSV — a stable Power BI source that,
    unlike an open .xlsx, is never locked by someone viewing it. Returns the dir."""
    if not os.path.isfile(master_path):
        return None
    csv_dir = csv_dir or os.path.join(os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(master_path)))), 'data', 'powerbi'), 'TUA')
    os.makedirs(csv_dir, exist_ok=True)
    try:
        sheets = pd.read_excel(master_path, sheet_name=None)
    except Exception:
        return None
    # Enrich with the static lookup tables (Critério, AdvT, Frequencies, Requests)
    # so the CSVs are dashboard-ready; the master workbook itself stays raw.
    try:
        import enrich
        static = os.getenv('EPAL_STATIC_TABLES',
                           os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(master_path)))),
                                        'data', 'powerbi', 'static_tables.xlsx'))
        sheets, _ = enrich.enrich_all(sheets, static)
    except Exception:
        pass                                          # never let enrichment break the export
    for name, df in sheets.items():
        df.to_csv(os.path.join(csv_dir, name + '.csv'), index=False, encoding='utf-8-sig')
    return csv_dir


def resolve_pdf_path(argv):
    if len(argv) > 1:
        return argv[1]
    env = os.getenv('EPAL_TUA_PDF_FILE')
    if env:
        return env
    # fall back to the first PDF in ./PDF files
    here = os.path.dirname(os.path.abspath(__file__))
    pdf_dir = os.path.join(here, 'PDF files')
    if os.path.isdir(pdf_dir):
        for f in sorted(os.listdir(pdf_dir)):
            if f.lower().endswith('.pdf'):
                return os.path.join(pdf_dir, f)
    return None


def main(argv):
    pdf = resolve_pdf_path(argv)
    if not pdf or not os.path.isfile(pdf):
        print(f"ERROR: TUA PDF not found. Pass a path: python extract_tua.py <file.pdf>  (got: {pdf})")
        return 2
    out_dir = os.getenv('EPAL_TUA_OUTPUT_DIR', os.path.dirname(os.path.abspath(__file__)))
    data = extract(pdf)
    jp, xp = write_outputs(data, out_dir)
    print(f"Extracted {len(data['condicoes_rejeicao'])} discharge parameters and "
          f"{len(data['autocontrolo'])} autocontrolo rows from {data['file_name']}")
    if data['_warnings']:
        print("WARNINGS:", '; '.join(data['_warnings']))
    print("JSON :", jp)
    print("Excel:", xp)
    # Upsert into the consolidated multi-sheet master (the feed for Power BI)
    master_path = os.getenv('EPAL_TUA_MASTER', os.path.join(out_dir, 'master_tua.xlsx'))
    mp, counts, deferred = update_master(data, master_path)
    summary = ', '.join(f"{n}={counts[n]}" for n in MASTER_SHEETS)
    if deferred:
        print(f"WARNING: '{master_path}' is open/locked in Excel — the master was NOT updated.")
        print(f"         Saved to '{mp}' instead. Close the master and re-run to consolidate.")
        return 3
    print(f"Master: {mp}  ({summary})")
    # Consolidated all-sections workbook (one sheet per section, every certificate)
    sections_path = os.getenv('EPAL_TUA_SECTIONS_MASTER', os.path.join(out_dir, 'master_all_sections.xlsx'))
    sp, scounts, sdef = update_sections_master(data, sections_path)
    print(f"All-sections master: {sp}  ({len(scounts)} section sheets)"
          + ("  [locked — saved to PENDING]" if sdef else ""))
    csv_dir = export_master_csv(master_path)
    if csv_dir:
        print("Power BI CSVs:", csv_dir)
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv))
